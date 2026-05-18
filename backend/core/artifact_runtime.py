"""Artifact Runtime — deterministic execution of approved artifacts.

ParserArtifact: deterministic parsing of bank/manual Excel files into CANONICAL_12 rows.
RuleArtifact: deterministic template filling from approved rules.

Implementation status:
    run_parser  — implemented (ParserArtifact deterministic runtime).
    run_rule    — contract placeholder only (Phase H1).
"""
from __future__ import annotations

import traceback
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Iterator, Optional

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from sqlalchemy.orm import Session

from core.artifact_ast_guard import PrimitivesViolationError
from core.artifact_sandbox import get_default_sandbox_config
from core.runtime_guard import no_ai_runtime
from db.tables import ParserArtifact


# ── Structured exception hierarchy ──

class ArtifactRuntimeError(RuntimeError):
    """Base for all artifact runtime errors."""


class ArtifactNotFoundError(ArtifactRuntimeError):
    """Artifact ID does not exist."""


class ArtifactNotActiveError(ArtifactRuntimeError):
    """Artifact exists but status != 'active'."""


class SandboxTimeoutError(ArtifactRuntimeError):
    """Artifact execution exceeded the time limit."""

    def __init__(self, artifact_id: int, timeout_seconds: int) -> None:
        self.artifact_id = artifact_id
        self.timeout_seconds = timeout_seconds
        super().__init__(
            f"Artifact {artifact_id} execution timed out after {timeout_seconds}s"
        )


class ArtifactExecutionError(ArtifactRuntimeError):
    """Artifact code raised an error during execution."""

    def __init__(self, artifact_id: int, detail: str) -> None:
        self.artifact_id = artifact_id
        self.detail = detail
        super().__init__(f"Artifact {artifact_id} execution failed: {detail}")


_SAFE_BUILTINS = {
    k: v for k, v in __builtins__.items()
    if k not in ("open", "eval", "exec", "compile", "breakpoint",
                  "memoryview", "globals", "locals", "vars", "dir")
} if isinstance(__builtins__, dict) else {
    k: getattr(__builtins__, k)
    for k in dir(__builtins__)
    if k not in ("open", "eval", "exec", "compile", "breakpoint",
                 "memoryview", "globals", "locals", "vars", "dir")
    and not k.startswith("__")
}


def run_parser(
    db: Session,
    artifact_id: int,
    file_path: str,
    ctx: Optional[dict[str, Any]] = None,
) -> Iterator[dict[str, Any]]:
    """Execute an active ParserArtifact against an Excel file.

    Input contract:
        db          — SQLAlchemy Session (read/write)
        artifact_id — int, references a ParserArtifact with status='active'
        file_path   — str, absolute path to the Excel file to parse
        ctx         — optional dict with account_code, entity_code, etc.

    Output contract:
        Yields dicts, each conforming to CANONICAL_12 schema (§C1).

    Implementation: Phase E1 delivered deterministic runtime.

    Raises:
        ArtifactNotFoundError    — artifact_id does not exist
        ArtifactNotActiveError   — artifact exists but status != 'active'
        PrimitivesViolationError — AST scan finds disallowed imports
        ArtifactExecutionError   — artifact code raises during execution
    """
    if ctx is None:
        ctx = {}

    from core.artifact_ast_guard import validate_artifact_code

    artifact = db.query(ParserArtifact).filter(ParserArtifact.id == artifact_id).first()
    if artifact is None:
        raise ArtifactNotFoundError(f"ParserArtifact id={artifact_id} 不存在")
    if artifact.status != "active":
        raise ArtifactNotActiveError(
            f"ParserArtifact id={artifact_id} status={artifact.status!r}，需要 'active'"
        )

    validate_artifact_code(artifact.code, artifact.id)

    wb = load_workbook(file_path, data_only=True)
    sandbox = get_default_sandbox_config()

    ns: dict[str, Any] = {"__builtins__": _SAFE_BUILTINS}
    code_obj = compile(artifact.code, f"<parser_artifact_{artifact.id}>", "exec")

    with no_ai_runtime():
        exec(code_obj, ns)
        parse_fn = ns.get("parse")
        if parse_fn is None or not callable(parse_fn):
            raise ArtifactExecutionError(
                artifact.id,
                "artifact code must define a callable 'parse(wb, ctx)' function"
            )
        try:
            result = parse_fn(wb, ctx)
        except ArtifactRuntimeError:
            raise
        except Exception as e:
            raise ArtifactExecutionError(
                artifact.id, f"{type(e).__name__}: {e}\n{traceback.format_exc()}"
            ) from e

    if callable(getattr(result, "__next__", None)):
        iterator = result
    else:
        iterator = iter(result)

    count = 0
    while count < sandbox.max_output_rows:
        try:
            row = next(iterator)
        except StopIteration:
            break
        except ArtifactRuntimeError:
            raise
        except Exception as e:
            raise ArtifactExecutionError(
                artifact.id, f"row {count}: {type(e).__name__}: {e}"
            ) from e
        if not isinstance(row, dict):
            raise ArtifactExecutionError(
                artifact.id, f"row {count}: expected dict, got {type(row).__name__}"
            )
        count += 1
        yield _coerce_row(artifact.id, row, count - 1)


def _coerce_row(artifact_id: int, row: dict[str, Any], index: int) -> dict[str, Any]:
    out = dict(row)
    for key in ("amount_in", "amount_out", "rolling_balance"):
        v = out.get(key)
        if v is not None and not isinstance(v, Decimal):
            try:
                out[key] = Decimal(str(v))
            except Exception:
                pass
    bd = out.get("business_date")
    if isinstance(bd, datetime):
        out["business_date"] = bd.date()
    elif isinstance(bd, str) and bd:
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y年%m月%d日"):
            try:
                out["business_date"] = datetime.strptime(bd, fmt).date()
                break
            except ValueError:
                continue
    return out


def run_rule(
    db: Session,
    artifact_id: int,
    ctx: dict[str, Any],
) -> Workbook:
    """Execute an active RuleArtifact to fill a report template.

    Implementation status:
        Contract placeholder only. Phase H1 will deliver the first working executor.

    Raises:
        NotImplementedError always.
    """
    raise NotImplementedError(
        "RuleArtifact runtime 尚未实现。"
        "完整执行器将在后续 Phase H1 中交付。"
    )
